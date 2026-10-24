"""Turning a module's mock data records into JSON or a workbook.

Isolated from the service, matching `app/services/checklist_export.py`. Unlike the
checklist's fixed columns, a mock data record's field set is dynamic, so the workbook's
header is the union of every record's keys, in first-seen order -- a record missing a
key that another record introduced later gets a blank cell rather than failing.
"""

import json
import re
from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.language import Language
from app.models.mock_data import MockDataRecord
from app.services.export_labels import EXPORT_LABELS

NO_DATA_LABEL = EXPORT_LABELS[Language.EN].no_data

MAX_SHEET_TITLE = 31
_ILLEGAL_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def build_mock_data_json(records: list[MockDataRecord]) -> bytes:
    """An array of field maps, one per record, in the given order."""
    return json.dumps([record.fields for record in records], indent=2).encode()


def _union_keys(records: list[MockDataRecord]) -> list[str]:
    """Every field key across every record, first-seen order, no duplicates."""
    seen: dict[str, None] = {}
    for record in records:
        for key in record.fields:
            seen.setdefault(key, None)
    return list(seen)


def _write_records(sheet: Worksheet, records: list[MockDataRecord]) -> None:
    """One column per field key (union across records), header frozen, values wrapped."""
    columns = _union_keys(records)
    sheet.append(columns)
    for index in range(1, len(columns) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 24
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    for record in records:
        sheet.append([record.fields.get(column) for column in columns])

    # Apply text wrapping to data rows for readability of long field values
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _save(book: Workbook) -> bytes:
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def build_mock_data_workbook(records: list[MockDataRecord], *, language: Language) -> bytes:
    """One sheet, one column per field key (union across records), header frozen."""
    book = Workbook()
    sheet = book.worksheets[0]
    sheet.title = EXPORT_LABELS[language].mock_data_sheet
    _write_records(sheet, records)
    return _save(book)


@dataclass(frozen=True, slots=True)
class ModuleRecords:
    """One module's part of a multi-module export, resolved by the caller."""

    name: str
    records: list[MockDataRecord]
    stale: bool


def sheet_title(name: str, taken: set[str], *, fallback: str = "Module") -> str:
    """A legal, unique worksheet name for a module.

    Excel refuses `[]:*?/\\` in a sheet name, caps it at 31 characters, and compares
    names case-insensitively, so two modules whose names differ only past the cut, or
    only in case, still need two sheets. `taken` holds the lower-cased names already
    used and is updated in place. `fallback` names a sheet whose module name cleans
    to nothing.
    """
    cleaned = _ILLEGAL_SHEET_CHARS.sub("_", name).strip("'").strip() or fallback
    candidate = cleaned[:MAX_SHEET_TITLE]
    suffix = 2
    while candidate.lower() in taken:
        tail = f" ({suffix})"
        candidate = cleaned[: MAX_SHEET_TITLE - len(tail)] + tail
        suffix += 1
    taken.add(candidate.lower())
    return candidate


def build_project_mock_data_workbook(modules: list[ModuleRecords], *, language: Language) -> bytes:
    """A cover sheet naming every chosen module, then one sheet per module with records.

    A module with no records gets a cover row and no sheet. Each module keeps its own
    columns: two modules' records rarely share a field set, and one union across a
    whole project would be mostly blank cells. The cover says what the file covered,
    so a module left without a sheet is visible rather than silently missing.
    """
    labels = EXPORT_LABELS[language]
    book = Workbook()
    cover = book.worksheets[0]
    cover.title = labels.cover_sheet
    cover.append(list(labels.cover_headers))
    for index, width in enumerate((36, 10, 8, 32), start=1):
        cover.column_dimensions[get_column_letter(index)].width = width
    for cell in cover[1]:
        cell.font = Font(bold=True)
    cover.freeze_panes = "A2"

    taken = {labels.cover_sheet.lower()}
    for module in modules:
        if not module.records:
            cover.append([module.name, 0, "", labels.no_data])
            continue
        title = sheet_title(module.name, taken, fallback=labels.cover_headers[0])
        cover.append(
            [module.name, len(module.records), labels.yes if module.stale else labels.no, title]
        )
        _write_records(book.create_sheet(title), module.records)
    return _save(book)
