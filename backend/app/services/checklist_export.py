"""Workbook construction for the checklist export.

Isolated from the service so the service stays about business rules. `.xlsx` rather
than CSV was a deliberate choice, and the argument is unchanged from the QA List it
replaces: the long columns are multi-paragraph prose, and CSV renders them as one
unwrapped line that runs off the screen -- the point of exporting a checklist is that
a person reads it and works through it.
"""

import uuid
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.models.checklist import ChecklistItem

SHEET_TITLE = "QA Checklist"

# Column header, and the width it gets. The prose columns are wide and wrap.
COLUMNS: tuple[tuple[str, int], ...] = (
    ("Module", 18),
    ("Feature", 18),
    ("Test name", 50),
    ("Expected result", 60),
    ("Current result", 60),
    ("Status", 14),
    ("Notes", 40),
    ("Source", 12),
    ("Reviewed by", 22),
    ("Reviewed at", 20),
    ("Citations", 40),
    ("Project", 36),
)
WRAPPED_COLUMNS = frozenset(
    {"Test name", "Expected result", "Current result", "Notes", "Citations"}
)


def _citations(item: ChecklistItem) -> str:
    """`path:start-end` per line. Without these the sheet cannot be audited against
    the repository, which is most of why anyone exports it."""
    return "\n".join(
        f"{citation.get('file_path')}:{citation.get('start_line')}-{citation.get('end_line')}"
        for citation in item.citations or []
    )


def build_workbook(
    rows: list[ChecklistItem],
    *,
    names: dict[uuid.UUID, str],
    module_names: dict[uuid.UUID, str],
) -> bytes:
    """One sheet, header frozen, prose wrapped.

    `names` maps user and project ids to something a human recognises, and
    `module_names` does the same for modules. The caller resolves both in one query
    each rather than letting this function touch the database. `rows` is expected
    already sorted `(module, feature, position)` -- the repository's `list_all`
    returns them that way and this function does not re-sort.
    """
    book = Workbook()
    # `book.active` is typed `Worksheet | None` because a workbook loaded from disk
    # could have zero sheets; a freshly constructed one always has exactly one, so
    # `worksheets[0]` gets the same sheet with a type mypy accepts unconditionally.
    sheet = book.worksheets[0]
    sheet.title = SHEET_TITLE

    sheet.append([header for header, _ in COLUMNS])
    for index, (_, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    # So the header stays visible while someone works down a hundred test cases.
    sheet.freeze_panes = "A2"

    for item in rows:
        sheet.append(
            [
                module_names.get(item.module_id, str(item.module_id)),
                item.feature,
                item.test_name,
                item.expected_result,
                item.current_result,
                item.status,
                item.notes,
                item.source,
                names.get(item.reviewed_by, "") if item.reviewed_by else "",
                item.reviewed_at.replace(tzinfo=None) if item.reviewed_at else None,
                _citations(item),
                names.get(item.project_id, str(item.project_id)),
            ]
        )

    wrapped = {
        index for index, (header, _) in enumerate(COLUMNS, start=1) if header in WRAPPED_COLUMNS
    }
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column in wrapped:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
