"""Workbook construction for the QA List export.

Isolated from the service so the service stays about business rules. `.xlsx` rather
than CSV was a deliberate choice (spec §6): the three long columns are
multi-paragraph prose, and CSV renders them as one unwrapped line that runs off the
screen — the point of exporting a regression set is that a person reads it.
"""

import uuid
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.models.qa_pair import QAPair

SHEET_TITLE = "QA Pairs"

# Column header, and the width it gets. The three prose columns are wide and wrap.
COLUMNS: tuple[tuple[str, int], ...] = (
    ("Module", 18),
    ("Question", 50),
    ("Expected result", 60),
    ("Result", 60),
    ("Status", 14),
    ("Tags", 20),
    ("Source", 12),
    ("Project", 36),
    ("Created by", 22),
    ("Reviewed by", 22),
    ("Last run", 20),
    ("Created at", 20),
    ("Citations", 40),
)
WRAPPED_COLUMNS = frozenset({"Question", "Expected result", "Result", "Citations"})


def _citations(pair: QAPair) -> str:
    """`path:start-end` per line. Without these the sheet cannot be audited
    against the repository, which is most of why anyone exports it."""
    return "\n".join(
        f"{citation.get('file_path')}:{citation.get('start_line')}-{citation.get('end_line')}"
        for citation in pair.citations or []
    )


def build_workbook(rows: list[QAPair], *, names: dict[uuid.UUID, str]) -> bytes:
    """One sheet, header frozen, prose wrapped.

    `names` maps user and project ids to something a human recognises. The service
    resolves it in one query rather than letting this function touch the database.
    """
    book = Workbook()
    # `book.active` is typed `Worksheet | None` because a workbook loaded from disk
    # could have zero sheets; a freshly constructed one always has exactly one, so
    # `worksheets[0]` gets the same sheet with a type mypy accepts unconditionally.
    sheet = book.worksheets[0]
    sheet.title = SHEET_TITLE

    sheet.append([header for header, _ in COLUMNS])
    for index, (_, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    # So the header stays visible while someone scrolls a hundred pairs.
    sheet.freeze_panes = "A2"

    for pair in rows:
        sheet.append(
            [
                pair.module,
                pair.question,
                pair.reference_answer,
                pair.answer,
                pair.status,
                ", ".join(pair.tags),
                pair.source,
                names.get(pair.project_id, str(pair.project_id)),
                names.get(pair.created_by, str(pair.created_by)),
                names.get(pair.reviewed_by, "") if pair.reviewed_by else "",
                pair.last_run_at.replace(tzinfo=None) if pair.last_run_at else None,
                pair.created_at.replace(tzinfo=None),
                _citations(pair),
            ]
        )

    wrapped = {
        index for index, (header, _) in enumerate(COLUMNS, start=1) if header in WRAPPED_COLUMNS
    }
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column in wrapped:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
