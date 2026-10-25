"""`GET /projects/{id}/export`: a whole project, or chosen modules, in one download."""

import dataclasses
import uuid
import zipfile
from io import BytesIO

import pytest
from fastapi import status
from httpx import AsyncClient, Response
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import Settings
from app.core.audit import AuditEventType, AuditRecorder
from app.core.errors import AppError, ErrorCode
from app.core.language import Language
from app.core.middleware import ProjectGrant
from app.core.permissions import EDITOR_NAME, VIEWER_NAME, Permission
from app.db.session import get_sessionmaker
from app.models.checklist import ChecklistItemStatus
from app.models.mock_data import MockDataDataset, MockDataDatasetStatus
from app.models.project import Project
from app.models.user import User
from app.schemas.project_export import ProjectExportPart, ProjectExportQuery
from app.services.export_labels import EXPORT_LABELS
from app.services.mock_data_export import sheet_title
from app.services.project_export import ProjectExportService
from tests.conftest import AuditRows, GrantMembership
from tests.factories import (
    create_checklist_item,
    create_checklist_module,
    create_mock_data_record,
    create_project,
    create_user,
)
from tests.helpers import authenticated


async def _project_with_modules(session: AsyncSession) -> tuple[Project, dict[str, uuid.UUID]]:
    """A project with three modules: Billing and Auth hold data, Search holds none.

    Created out of name order, so the export's by-name ordering is what is tested.
    """
    project = await create_project(session, name="api")
    project.active_generation = 2
    ids: dict[str, uuid.UUID] = {}
    for name in ("Billing", "Search", "Auth"):
        module = await create_checklist_module(
            session, project_id=project.id, created_by=project.created_by, name=name
        )
        ids[name] = module.id
    for name, tests in (("Billing", ["Refunds a charge"]), ("Auth", ["Rejects a bad password"])):
        for position, test_name in enumerate(tests):
            await create_checklist_item(
                session,
                module_id=ids[name],
                project_id=project.id,
                created_by=project.created_by,
                test_name=test_name,
                position=position,
            )
    await create_checklist_item(
        session,
        module_id=ids["Auth"],
        project_id=project.id,
        created_by=project.created_by,
        test_name="Locks after five tries",
        status=ChecklistItemStatus.FAIL,
        position=1,
    )
    await create_mock_data_record(
        session, module_id=ids["Auth"], fields={"email": "a@example.com", "role": "admin"}
    )
    await create_mock_data_record(session, module_id=ids["Billing"], fields={"amount": "12.50"})
    # Billing's data was generated against generation 1; the project is on 2.
    session.add(
        MockDataDataset(
            id=uuid.uuid4(),
            checklist_module_id=ids["Billing"],
            status=MockDataDatasetStatus.READY.value,
            indexed_generation=1,
        )
    )
    await session.commit()
    return project, ids


def _sheet_rows(content: bytes, title: str | None = None) -> list[tuple[object, ...]]:
    book = load_workbook(BytesIO(content))
    sheet = book[title] if title else book.worksheets[0]
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


async def _export(client: AsyncClient, project_id: uuid.UUID, query: str) -> Response:
    return await client.get(f"/projects/{project_id}/export?{query}")


async def test_the_checklist_part_covers_every_module_grouped_by_name(
    authed_client: AsyncClient,
    authed_user: User,
    grant_membership: GrantMembership,
    db_session: AsyncSession,
    audit_rows: AuditRows,
) -> None:
    project, _ = await _project_with_modules(db_session)
    await grant_membership(authed_user.id, project.id, VIEWER_NAME)

    response = await _export(authed_client, project.id, "include=checklist")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="qa-checklist.xlsx"' in response.headers["content-disposition"]
    header, *rows = _sheet_rows(response.content)
    assert header[0] == "Module"
    assert [(row[0], row[2]) for row in rows] == [
        ("Auth", "Rejects a bad password"),
        ("Auth", "Locks after five tries"),
        ("Billing", "Refunds a charge"),
    ]
    (event,) = await audit_rows(AuditEventType.CHECKLIST_EXPORTED)
    assert (event.target_type, event.target_id, event.project_id) == (
        "project",
        project.id,
        project.id,
    )
    assert event.details == {
        "format": "xlsx",
        "rowCount": 3,
        "filter": {"projectId": str(project.id)},
        "projectCount": 1,
        "moduleCount": 3,
    }
    assert "Refunds" not in str(event.details)


async def test_the_checklist_part_uses_the_projects_content_language(
    authed_client: AsyncClient,
    authed_user: User,
    grant_membership: GrantMembership,
    db_session: AsyncSession,
) -> None:
    project, _ = await _project_with_modules(db_session)
    project.content_language = "id"
    await db_session.commit()
    await grant_membership(authed_user.id, project.id, VIEWER_NAME)

    response = await _export(authed_client, project.id, "include=checklist")

    header, *_ = _sheet_rows(response.content)
    assert list(header) == list(EXPORT_LABELS[Language.ID].checklist_headers)


async def test_the_grid_filters_narrow_the_checklist_part(
    authed_client: AsyncClient,
    authed_user: User,
    grant_membership: GrantMembership,
    db_session: AsyncSession,
    audit_rows: AuditRows,
) -> None:
    project, _ = await _project_with_modules(db_session)
    await grant_membership(authed_user.id, project.id, VIEWER_NAME)

    response = await _export(authed_client, project.id, "include=checklist&status=fail&feature=Log")

    assert response.status_code == 200
    _, *rows = _sheet_rows(response.content)
    assert [row[2] for row in rows] == ["Locks after five tries"]
    (event,) = await audit_rows(AuditEventType.CHECKLIST_EXPORTED)
    assert event.details["filter"] == {
        "projectId": str(project.id),
        "featureFilterApplied": "true",
        "status": "fail",
    }


async def test_the_mock_data_part_lists_every_module_on_a_cover_sheet(
    authed_client: AsyncClient,
    authed_user: User,
    grant_membership: GrantMembership,
    db_session: AsyncSession,
    audit_rows: AuditRows,
) -> None:
    project, _ = await _project_with_modules(db_session)
    await grant_membership(authed_user.id, project.id, VIEWER_NAME)

    response = await _export(authed_client, project.id, "include=mockData")

    assert response.status_code == 200
    assert 'filename="mock-data.xlsx"' in response.headers["content-disposition"]
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Modules", "Auth", "Billing"]
    assert _sheet_rows(response.content, "Modules") == [
        ("Module", "Records", "Stale", "Sheet"),
        ("Auth", 1, "no", "Auth"),
        ("Billing", 1, "yes", "Billing"),
        ("Search", 0, None, EXPORT_LABELS[Language.EN].no_data),
    ]
    # Each module keeps its own columns rather than one union across the project.
    # Compared as a mapping: `jsonb` does not keep the order keys were written in.
    header, values = _sheet_rows(response.content, "Auth")
    assert dict(zip(header, values, strict=True)) == {"email": "a@example.com", "role": "admin"}
    (event,) = await audit_rows(AuditEventType.MOCK_DATA_EXPORTED)
    assert event.target_type == "project"
    assert event.details == {"format": "xlsx", "rowCount": 2, "moduleCount": 3}


async def test_both_parts_arrive_as_one_zip_and_audit_once_each(
    authed_client: AsyncClient,
    authed_user: User,
    grant_membership: GrantMembership,
    db_session: AsyncSession,
    audit_rows: AuditRows,
) -> None:
    project, _ = await _project_with_modules(db_session)
    await grant_membership(authed_user.id, project.id, VIEWER_NAME)

    response = await _export(authed_client, project.id, "include=checklist&include=mockData")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    assert 'filename="qa-export.zip"' in response.headers["content-disposition"]
    with zipfile.ZipFile(BytesIO(response.content)) as archive:
        assert sorted(archive.namelist()) == ["mock-data.xlsx", "qa-checklist.xlsx"]
        assert _sheet_rows(archive.read("qa-checklist.xlsx"))[0][0] == "Module"
    assert len(await audit_rows(AuditEventType.CHECKLIST_EXPORTED)) == 1
    assert len(await audit_rows(AuditEventType.MOCK_DATA_EXPORTED)) == 1


async def test_chosen_modules_narrow_both_parts(
    authed_client: AsyncClient,
    authed_user: User,
    grant_membership: GrantMembership,
    db_session: AsyncSession,
    audit_rows: AuditRows,
) -> None:
    project, ids = await _project_with_modules(db_session)
    await grant_membership(authed_user.id, project.id, VIEWER_NAME)

    response = await _export(
        authed_client,
        project.id,
        f"include=checklist&include=mockData&moduleIds={ids['Billing']}&moduleIds={ids['Search']}",
    )

    assert response.status_code == 200
    with zipfile.ZipFile(BytesIO(response.content)) as archive:
        _, *checklist = _sheet_rows(archive.read("qa-checklist.xlsx"))
        cover = _sheet_rows(archive.read("mock-data.xlsx"), "Modules")
    assert [row[0] for row in checklist] == ["Billing"]
    assert [row[0] for row in cover[1:]] == ["Billing", "Search"]
    (event,) = await audit_rows(AuditEventType.MOCK_DATA_EXPORTED)
    assert event.details["moduleCount"] == 2


async def test_another_projects_module_is_a_404_and_records_nothing(
    authed_client: AsyncClient,
    authed_user: User,
    grant_membership: GrantMembership,
    db_session: AsyncSession,
    audit_rows: AuditRows,
) -> None:
    """Even one the caller can read elsewhere: the id is not *this* project's module."""
    project, ids = await _project_with_modules(db_session)
    elsewhere = await create_checklist_module(db_session, name="Elsewhere")
    await db_session.commit()
    await grant_membership(authed_user.id, project.id, VIEWER_NAME)
    await grant_membership(authed_user.id, elsewhere.project_id, EDITOR_NAME)

    response = await _export(
        authed_client,
        project.id,
        f"include=checklist&moduleIds={ids['Auth']}&moduleIds={elsewhere.id}",
    )

    assert response.status_code == 404
    assert response.json()["detail"]["code"] == "CHECKLIST_MODULE_NOT_FOUND"
    assert await audit_rows(AuditEventType.CHECKLIST_EXPORTED) == []


async def test_a_non_member_cannot_learn_the_project_exists(
    authed_client: AsyncClient, db_session: AsyncSession
) -> None:
    project, _ = await _project_with_modules(db_session)

    response = await _export(authed_client, project.id, "include=checklist")

    assert response.status_code == 404
    assert response.json()["detail"]["code"] == "PROJECT_NOT_FOUND"


@pytest.mark.parametrize("query", ["", "include=everything", "include=checklist&status=bogus"])
async def test_a_malformed_query_is_a_422(
    authed_client: AsyncClient,
    authed_user: User,
    grant_membership: GrantMembership,
    db_session: AsyncSession,
    query: str,
) -> None:
    project, _ = await _project_with_modules(db_session)
    await grant_membership(authed_user.id, project.id, VIEWER_NAME)

    assert (await _export(authed_client, project.id, query)).status_code == 422


def _service(session: AsyncSession, settings: Settings | None = None) -> ProjectExportService:
    return ProjectExportService(
        session, settings or Settings(), recorder=AuditRecorder(get_sessionmaker())
    )


async def test_a_member_who_cannot_read_mock_data_gets_a_403(
    db_session: AsyncSession, grant_membership: GrantMembership
) -> None:
    """A custom role may hold `checklist.read` without `mockdata.read`; each part is
    gated on its own permission, and the member can see the project, so `403`."""
    project, _ = await _project_with_modules(db_session)
    member = await create_user(db_session)
    await grant_membership(member.id, project.id, VIEWER_NAME)
    actor = await authenticated(db_session, member)
    narrowed = dataclasses.replace(
        actor,
        grants={
            project.id: ProjectGrant(
                project_id=project.id,
                role="checklist-only",
                permissions=frozenset({Permission.PROJECT_READ, Permission.CHECKLIST_READ}),
            )
        },
    )

    with pytest.raises(AppError) as caught:
        await _service(db_session).export(
            project.id,
            ProjectExportQuery(include=[ProjectExportPart.CHECKLIST, ProjectExportPart.MOCK_DATA]),
            actor=narrowed,
        )

    assert caught.value.status_code == status.HTTP_403_FORBIDDEN
    assert caught.value.code is ErrorCode.INSUFFICIENT_ROLE


async def test_a_part_over_its_cap_refuses_the_whole_export_and_records_nothing(
    db_session: AsyncSession, grant_membership: GrantMembership, audit_rows: AuditRows
) -> None:
    """The checklist part fits; the mock data part does not. Nothing left, so the
    checklist part must not be recorded as exported either."""
    project, _ = await _project_with_modules(db_session)
    member = await create_user(db_session)
    await grant_membership(member.id, project.id, VIEWER_NAME)

    with pytest.raises(AppError) as caught:
        await _service(db_session, Settings(mock_data_export_max_rows=1)).export(
            project.id,
            ProjectExportQuery(include=[ProjectExportPart.CHECKLIST, ProjectExportPart.MOCK_DATA]),
            actor=await authenticated(db_session, member),
        )

    assert caught.value.code is ErrorCode.EXPORT_TOO_LARGE
    assert await audit_rows(AuditEventType.CHECKLIST_EXPORTED) == []
    assert await audit_rows(AuditEventType.MOCK_DATA_EXPORTED) == []


def test_sheet_titles_are_legal_and_unique() -> None:
    taken = {"modules"}
    titles = [
        sheet_title(name, taken)
        for name in ("Auth/Login [v2]", "auth/login [v2]", "x" * 40, "x" * 40, "Modules")
    ]

    assert titles == [
        "Auth_Login _v2_",
        "auth_login _v2_ (2)",
        "x" * 31,
        "x" * 27 + " (2)",
        "Modules (2)",
    ]
    assert all(len(title) <= 31 for title in titles)
