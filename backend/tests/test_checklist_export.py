"""Workbook construction for the checklist export."""

from io import BytesIO

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.language import Language
from app.models.checklist import ChecklistItemStatus
from app.models.project import Project
from app.services.checklist_export import COLUMNS, SHEET_TITLE, build_workbook
from app.services.export_labels import EXPORT_LABELS
from tests.factories import create_checklist_item, create_checklist_module


async def test_columns_and_order_match_the_grid(db_session: AsyncSession) -> None:
    """Sorted by (module, feature, position) so the sheet reads in the order a tester
    works -- which is most of why anyone exports it (spec 7)."""
    module = await create_checklist_module(db_session, name="Auth")
    rows = [
        await create_checklist_item(
            db_session,
            module_id=module.id,
            project_id=module.project_id,
            created_by=module.created_by,
            feature="Login",
            test_name=name,
            position=position,
            status=ChecklistItemStatus.PASS,
            current_result="Observed 401",
        )
        for position, name in ((0, "first"), (1, "second"))
    ]

    content = build_workbook(
        rows,
        names={module.project_id: "repo", module.created_by: "Test User"},
        module_names={module.id: "Auth"},
        language=Language.EN,
    )

    sheet = load_workbook(BytesIO(content)).worksheets[0]
    assert sheet.title == SHEET_TITLE
    assert [cell.value for cell in sheet[1]] == [header for header, _ in COLUMNS]
    assert sheet.cell(row=2, column=1).value == "Auth"
    assert sheet.cell(row=2, column=3).value == "first"
    assert sheet.cell(row=3, column=3).value == "second"
    assert sheet.freeze_panes == "A2"


async def test_an_untested_row_exports_an_empty_result(db_session: AsyncSession) -> None:
    """`current_result` is a human's observation, so an ungraded row is blank rather
    than filled with a prediction (spec 2.3)."""
    module = await create_checklist_module(db_session)
    item = await create_checklist_item(
        db_session,
        module_id=module.id,
        project_id=module.project_id,
        created_by=module.created_by,
    )

    sheet = load_workbook(
        BytesIO(
            build_workbook(
                [item],
                names={module.project_id: "repo", module.created_by: "u"},
                module_names={module.id: module.name},
                language=Language.EN,
            )
        )
    ).worksheets[0]

    headers = [header for header, _ in COLUMNS]
    assert sheet.cell(row=2, column=headers.index("Current result") + 1).value is None
    assert sheet.cell(row=2, column=headers.index("Status") + 1).value == "untested"


async def _export_headers(client: AsyncClient) -> list[object]:
    response = await client.get("/checklist-items/export")
    assert response.status_code == 200
    return [cell.value for cell in load_workbook(BytesIO(response.content)).worksheets[0][1]]


async def _module_in_language(db_session: AsyncSession, language: str) -> None:
    module = await create_checklist_module(db_session)
    await create_checklist_item(db_session, module_id=module.id, project_id=module.project_id)
    project = await db_session.get(Project, module.project_id)
    assert project is not None
    project.content_language = language


async def test_a_cross_project_export_in_one_shared_language_uses_it(
    client_for_admin: AsyncClient, db_session: AsyncSession
) -> None:
    await _module_in_language(db_session, "ms")
    await _module_in_language(db_session, "ms")
    await db_session.commit()

    assert await _export_headers(client_for_admin) == list(
        EXPORT_LABELS[Language.MS].checklist_headers
    )


async def test_a_cross_project_export_in_mixed_languages_is_english(
    client_for_admin: AsyncClient, db_session: AsyncSession
) -> None:
    await _module_in_language(db_session, "ms")
    await _module_in_language(db_session, "ja")
    await db_session.commit()

    assert await _export_headers(client_for_admin) == list(
        EXPORT_LABELS[Language.EN].checklist_headers
    )
