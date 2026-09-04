"""Item business rules. The two writes, and the split between them, are the point."""

import uuid

import pytest
from fastapi import status
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import Settings
from app.core.errors import AppError, ErrorCode
from app.models.checklist import ChecklistItemSource, ChecklistItemStatus
from app.schemas.checklist import (
    ChecklistItemCreateRequest,
    ChecklistItemListQuery,
    ChecklistItemResultRequest,
    ChecklistItemUpdateRequest,
    ChecklistResultsClearRequest,
)
from app.services.checklist_item import ChecklistItemService
from tests.factories import create_checklist_item, create_checklist_module, create_user
from tests.helpers import authenticated


async def test_any_user_may_record_a_result(db_session: AsyncSession) -> None:
    """A tester who did not author the checklist must be able to record what they
    observed. Ungated, deliberately (spec 2.5)."""
    module = await create_checklist_module(db_session)
    item = await create_checklist_item(
        db_session,
        module_id=module.id,
        project_id=module.project_id,
        created_by=module.created_by,
    )
    tester = await create_user(db_session)
    service = ChecklistItemService(db_session, Settings())

    result = await service.set_result(
        item.id,
        ChecklistItemResultRequest(current_result="Returned 500", status=ChecklistItemStatus.FAIL),
        actor=authenticated(tester),
    )

    assert result.current_result == "Returned 500"
    assert result.status is ChecklistItemStatus.FAIL
    # Who looked, and when. Without it a verdict is an anonymous claim.
    assert result.reviewed_by == tester.id
    assert result.reviewed_at is not None


async def test_a_tester_cannot_rewrite_the_expectation(db_session: AsyncSession) -> None:
    """Otherwise the cheapest way to make a failing test pass is to edit what was
    expected -- which is the whole reason the two writes are separate routes."""
    module = await create_checklist_module(db_session)
    item = await create_checklist_item(
        db_session,
        module_id=module.id,
        project_id=module.project_id,
        created_by=module.created_by,
    )
    tester = await create_user(db_session)
    service = ChecklistItemService(db_session, Settings())

    with pytest.raises(AppError) as caught:
        await service.update(
            item.id,
            ChecklistItemUpdateRequest(expected_result="Anything is fine"),
            actor=authenticated(tester),
        )

    assert caught.value.status_code == status.HTTP_403_FORBIDDEN
    assert caught.value.code is ErrorCode.NOT_CHECKLIST_OWNER


async def test_setting_a_result_back_to_untested_clears_the_reviewer(
    db_session: AsyncSession,
) -> None:
    """`untested` means nobody has looked. Leaving a reviewer on it would say someone
    did, and the pass rate would be computed against a verdict that was withdrawn."""
    module = await create_checklist_module(db_session)
    item = await create_checklist_item(
        db_session,
        module_id=module.id,
        project_id=module.project_id,
        created_by=module.created_by,
        status=ChecklistItemStatus.PASS,
    )
    service = ChecklistItemService(db_session, Settings())
    tester = await create_user(db_session)
    await service.set_result(
        item.id,
        ChecklistItemResultRequest(current_result="ok", status=ChecklistItemStatus.PASS),
        actor=authenticated(tester),
    )

    result = await service.set_result(
        item.id,
        ChecklistItemResultRequest(current_result=None, status=ChecklistItemStatus.UNTESTED),
        actor=authenticated(tester),
    )

    assert result.reviewed_by is None
    assert result.reviewed_at is None
    assert result.current_result is None


async def test_a_manual_item_is_marked_manual_and_positioned_last(
    db_session: AsyncSession,
) -> None:
    module = await create_checklist_module(db_session)
    await create_checklist_item(
        db_session,
        module_id=module.id,
        project_id=module.project_id,
        created_by=module.created_by,
        feature="Login",
        position=0,
    )
    service = ChecklistItemService(db_session, Settings())

    created = await service.create(
        ChecklistItemCreateRequest(
            module_id=module.id,
            feature="Login",
            test_name="Rejects an empty password",
            expected_result="422 VALIDATION_ERROR",
        ),
        actor=authenticated(await create_user(db_session)),
    )

    assert created.source is ChecklistItemSource.MANUAL
    assert created.position == 1
    # Denormalised at insert and never updated: a module cannot move projects.
    assert created.project_id == module.project_id
    assert created.status is ChecklistItemStatus.UNTESTED
    assert created.current_result is None


async def test_export_refuses_above_the_cap(db_session: AsyncSession) -> None:
    """`openpyxl` allocates the whole book in memory even write-only, so the cap is
    the only thing bounding it."""
    module = await create_checklist_module(db_session)
    for index in range(3):
        await create_checklist_item(
            db_session,
            module_id=module.id,
            project_id=module.project_id,
            created_by=module.created_by,
            position=index,
        )
    service = ChecklistItemService(db_session, Settings(checklist_export_max_rows=2))

    with pytest.raises(AppError) as caught:
        await service.export(
            ChecklistItemListQuery(), actor=authenticated(await create_user(db_session))
        )

    assert caught.value.status_code == status.HTTP_409_CONFLICT
    assert caught.value.code is ErrorCode.EXPORT_TOO_LARGE


async def test_export_applies_the_same_filters_and_ignores_pagination(
    db_session: AsyncSession,
) -> None:
    """The point is to get the whole filtered set into one file (spec 7)."""
    module = await create_checklist_module(db_session)
    for index in range(3):
        await create_checklist_item(
            db_session,
            module_id=module.id,
            project_id=module.project_id,
            created_by=module.created_by,
            position=index,
            status=ChecklistItemStatus.PASS if index else ChecklistItemStatus.FAIL,
        )
    service = ChecklistItemService(db_session, Settings())

    content = await service.export(
        ChecklistItemListQuery(limit=1, status=ChecklistItemStatus.PASS),
        actor=authenticated(await create_user(db_session)),
    )

    from io import BytesIO

    from openpyxl import load_workbook

    sheet = load_workbook(BytesIO(content)).worksheets[0]
    assert sheet.max_row == 3  # header plus the two PASS rows, `limit` ignored


async def test_update_changes_the_definition_and_leaves_the_result_alone(
    db_session: AsyncSession,
) -> None:
    """The gated write's success path.

    Asserts the four definition fields change AND that `current_result`/`status`
    survive untouched -- the whole reason `ChecklistItemUpdateRequest` carries no
    field for them is that editing an expectation must never quietly discard a
    tester's recorded observation.
    """
    creator = await create_user(db_session)
    module = await create_checklist_module(db_session, created_by=creator.id)
    item = await create_checklist_item(
        db_session,
        module_id=module.id,
        project_id=module.project_id,
        created_by=creator.id,
    )
    service = ChecklistItemService(db_session, Settings())
    # A genuine recorded observation, not the factory's default -- via `set_result`,
    # the only path that actually writes these columns.
    await service.set_result(
        item.id,
        ChecklistItemResultRequest(current_result="Returned 500", status=ChecklistItemStatus.FAIL),
        actor=authenticated(creator),
    )

    updated = await service.update(
        item.id,
        ChecklistItemUpdateRequest(
            feature="Signup",
            test_name="Rejects a duplicate email",
            expected_result="409 EMAIL_ALREADY_EXISTS",
            notes="Edge case",
        ),
        actor=authenticated(creator),
    )

    assert updated.feature == "Signup"
    assert updated.test_name == "Rejects a duplicate email"
    assert updated.expected_result == "409 EMAIL_ALREADY_EXISTS"
    assert updated.notes == "Edge case"
    # The observation survives an edit to what was expected.
    assert updated.current_result == "Returned 500"
    assert updated.status is ChecklistItemStatus.FAIL


async def test_list_is_scoped_and_filters_within_the_scope(
    db_session: AsyncSession,
) -> None:
    """`list` returns items from modules the caller did not create -- phase 1 shares
    everything -- and its `module_id` filter narrows within that scope rather than
    replacing it."""
    creator = await create_user(db_session)
    module = await create_checklist_module(db_session, created_by=creator.id)
    item_in_module = await create_checklist_item(
        db_session,
        module_id=module.id,
        project_id=module.project_id,
        created_by=creator.id,
        feature="Login",
    )
    other_module = await create_checklist_module(db_session, created_by=creator.id)
    item_in_other_module = await create_checklist_item(
        db_session,
        module_id=other_module.id,
        project_id=other_module.project_id,
        created_by=creator.id,
        feature="Signup",
    )
    caller = await create_user(db_session)
    service = ChecklistItemService(db_session, Settings())

    unfiltered = await service.list(ChecklistItemListQuery(), actor=authenticated(caller))
    unfiltered_ids = {row.id for row in unfiltered.items}
    assert item_in_module.id in unfiltered_ids
    assert item_in_other_module.id in unfiltered_ids

    filtered = await service.list(
        ChecklistItemListQuery(module_id=module.id), actor=authenticated(caller)
    )
    filtered_ids = {row.id for row in filtered.items}
    assert filtered_ids == {item_in_module.id}


async def test_delete_is_gated_and_hides_the_item(db_session: AsyncSession) -> None:
    """Soft delete: a stranger gets 403 (existence is not a secret), the creator
    succeeds, and the item stops appearing in `list` afterwards."""
    creator = await create_user(db_session)
    module = await create_checklist_module(db_session, created_by=creator.id)
    item = await create_checklist_item(
        db_session,
        module_id=module.id,
        project_id=module.project_id,
        created_by=creator.id,
    )
    stranger = await create_user(db_session)
    service = ChecklistItemService(db_session, Settings())

    with pytest.raises(AppError) as caught:
        await service.delete(item.id, actor=authenticated(stranger))
    assert caught.value.status_code == status.HTTP_403_FORBIDDEN
    assert caught.value.code is ErrorCode.NOT_CHECKLIST_OWNER

    await service.delete(item.id, actor=authenticated(creator))

    remaining = await service.list(
        ChecklistItemListQuery(module_id=module.id), actor=authenticated(creator)
    )
    assert item.id not in {row.id for row in remaining.items}


async def test_clear_results_is_open_to_a_user_who_did_not_create_the_checklist(
    db_session: AsyncSession,
) -> None:
    """Ungated, exactly like the result write it undoes (spec 2.5).

    Gating this on `created_by` would let a tester record an observation and then not
    be allowed to take it back, which is the mirror image of the reason recording is
    open in the first place.
    """
    module = await create_checklist_module(db_session)
    item = await create_checklist_item(
        db_session,
        module_id=module.id,
        project_id=module.project_id,
        created_by=module.created_by,
        status=ChecklistItemStatus.PASS,
        current_result="200 OK",
    )
    service = ChecklistItemService(db_session, Settings())
    stranger = authenticated(await create_user(db_session))

    response = await service.clear_results(
        ChecklistResultsClearRequest(module_id=module.id), actor=stranger
    )

    assert response.cleared_count == 1
    await db_session.refresh(item)
    assert item.status == ChecklistItemStatus.UNTESTED.value
    assert item.current_result is None


async def test_clear_results_404s_on_a_module_outside_the_scope(
    db_session: AsyncSession,
) -> None:
    """Bounded rather than gated: the request names one module and the module has to
    be readable, so nothing can widen the blast radius past it."""
    service = ChecklistItemService(db_session, Settings())

    with pytest.raises(AppError) as caught:
        await service.clear_results(
            ChecklistResultsClearRequest(module_id=uuid.uuid4()),
            actor=authenticated(await create_user(db_session)),
        )

    assert caught.value.status_code == 404
