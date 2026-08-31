"""The spreadsheet export: filters, the cap, the headers, and a readable book."""

import io

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from tests.factories import create_project, create_qa_pair, create_user

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def test_export_returns_a_workbook_with_the_expected_header_row(
    client_for_user_a: AsyncClient, db_session: AsyncSession
) -> None:
    user = await create_user(db_session)
    await create_qa_pair(db_session, created_by=user.id, module="auth", tags=["auth"])
    await db_session.commit()

    response = await client_for_user_a.get("/qa-pairs/export")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
    assert "attachment" in response.headers["content-disposition"]
    assert ".xlsx" in response.headers["content-disposition"]

    book = load_workbook(io.BytesIO(response.content))
    sheet = book["QA Pairs"]
    headers = [cell.value for cell in sheet[1]]
    assert headers[:5] == ["Module", "Question", "Expected result", "Result", "Status"]
    assert "Citations" in headers
    assert sheet.max_row == 2


async def test_export_honours_the_same_filters_as_the_list(
    client_for_user_a: AsyncClient, db_session: AsyncSession
) -> None:
    user = await create_user(db_session)
    await create_qa_pair(db_session, created_by=user.id, tags=["auth"])
    await create_qa_pair(db_session, created_by=user.id, tags=["billing"])
    await db_session.commit()

    response = await client_for_user_a.get("/qa-pairs/export?tag=auth")

    book = load_workbook(io.BytesIO(response.content))
    # One header row plus one matching pair.
    assert book["QA Pairs"].max_row == 2


async def test_export_names_people_not_uuids(
    client_for_user_a: AsyncClient, db_session: AsyncSession
) -> None:
    """A spreadsheet of UUIDs is not readable by the person the export is for."""
    user = await create_user(db_session, name="Ada Lovelace")
    await create_qa_pair(db_session, created_by=user.id)
    await db_session.commit()

    response = await client_for_user_a.get("/qa-pairs/export")

    book = load_workbook(io.BytesIO(response.content))
    sheet = book["QA Pairs"]
    headers = [cell.value for cell in sheet[1]]
    created_by = sheet.cell(row=2, column=headers.index("Created by") + 1).value
    assert created_by == "Ada Lovelace"


async def test_export_refuses_more_rows_than_the_cap(
    client_for_user_a: AsyncClient, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    get_settings.cache_clear()
    monkeypatch.setenv("QA_EXPORT_MAX_ROWS", "1")

    user = await create_user(db_session)
    project = await create_project(db_session, created_by=user.id)
    await create_qa_pair(db_session, project_id=project.id, created_by=user.id)
    await create_qa_pair(db_session, project_id=project.id, created_by=user.id)
    await db_session.commit()

    response = await client_for_user_a.get("/qa-pairs/export")

    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "EXPORT_TOO_LARGE"
    get_settings.cache_clear()
