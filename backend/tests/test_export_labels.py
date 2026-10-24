"""Export labels: a fixed table per language, covering every column (spec §4.4)."""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.language import Language
from app.services.checklist_export import COLUMNS, build_workbook
from app.services.export_labels import EXPORT_LABELS
from app.services.mock_data_export import (
    ModuleRecords,
    build_mock_data_workbook,
    build_project_mock_data_workbook,
)

ILLEGAL = set("\\/*?:[]")


def test_every_language_has_labels() -> None:
    assert set(EXPORT_LABELS) == set(Language)


@pytest.mark.parametrize("language", list(Language))
def test_labels_fit_the_workbook(language: Language) -> None:
    labels = EXPORT_LABELS[language]

    assert len(labels.checklist_headers) == len(COLUMNS)
    for title in (labels.checklist_sheet, labels.mock_data_sheet, labels.cover_sheet):
        assert 0 < len(title) <= 31
        assert not ILLEGAL & set(title)


def test_english_headers_are_unchanged() -> None:
    """The English export is byte-for-byte what it was before this phase."""
    assert EXPORT_LABELS[Language.EN].checklist_headers == tuple(h for h, _ in COLUMNS)


def test_the_checklist_workbook_uses_the_languages_headers() -> None:
    content = build_workbook([], names={}, module_names={}, language=Language.ID)
    sheet = load_workbook(BytesIO(content)).worksheets[0]

    assert sheet.title == EXPORT_LABELS[Language.ID].checklist_sheet
    assert [cell.value for cell in sheet[1]] == list(EXPORT_LABELS[Language.ID].checklist_headers)


def test_the_mock_data_workbook_titles_its_sheet() -> None:
    content = build_mock_data_workbook([], language=Language.JA)

    sheet = load_workbook(BytesIO(content)).worksheets[0]
    assert sheet.title == EXPORT_LABELS[Language.JA].mock_data_sheet


def test_the_cover_sheet_is_translated() -> None:
    content = build_project_mock_data_workbook(
        [ModuleRecords(name="Login", records=[], stale=False)], language=Language.KO
    )
    cover = load_workbook(BytesIO(content)).worksheets[0]
    labels = EXPORT_LABELS[Language.KO]

    assert cover.title == labels.cover_sheet
    assert [cell.value for cell in cover[1]] == list(labels.cover_headers)
    assert cover["D2"].value == labels.no_data
