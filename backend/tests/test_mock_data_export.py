"""Export tests for mock data records.

Mirrors test_checklist_export.py's workbook-shape assertions, for the dynamic
field-map case, plus the JSON export and the row-cap refusal.
"""

import json
import uuid
from datetime import UTC, datetime, timedelta
from io import BytesIO

import pytest
from fastapi import status
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import Settings
from app.core.errors import AppError
from app.models.mock_data import MockDataRecord
from app.services.mock_data_dataset import MockDataDatasetService
from app.services.mock_data_export import build_mock_data_json, build_mock_data_workbook
from tests.factories import (
    create_checklist_module,
    create_mock_data_record,
    create_project,
    create_user,
)
from tests.helpers import authenticated


def test_build_mock_data_json_is_an_array_of_field_maps() -> None:
    """Pure function: JSON array of field maps, one per record."""
    record = MockDataRecord(
        id=uuid.uuid4(),
        checklist_module_id=uuid.uuid4(),
        fields={"name": "Acme"},
        created_by=uuid.uuid4(),
    )
    payload = json.loads(build_mock_data_json([record]))
    assert payload == [{"name": "Acme"}]


def test_build_mock_data_workbook_unions_field_keys_across_records() -> None:
    """Pure function: workbook with union of all field keys."""
    records = [
        MockDataRecord(
            id=uuid.uuid4(),
            checklist_module_id=uuid.uuid4(),
            fields={"name": "Acme"},
            created_by=uuid.uuid4(),
        ),
        MockDataRecord(
            id=uuid.uuid4(),
            checklist_module_id=uuid.uuid4(),
            fields={"name": "Globex", "start": "2026-01-01"},
            created_by=uuid.uuid4(),
        ),
    ]
    book = load_workbook(BytesIO(build_mock_data_workbook(records)))
    sheet = book.active
    header = [cell.value for cell in sheet[1]]
    assert header == ["name", "start"]
    assert [cell.value for cell in sheet[2]] == ["Acme", None]
    assert [cell.value for cell in sheet[3]] == ["Globex", "2026-01-01"]


async def test_export_json_succeeds_under_cap(db_session: AsyncSession) -> None:
    """export_json returns JSON array when records are under cap."""
    # Create project with embedding_collection
    project = await create_project(db_session)
    project.embedding_collection = "col"
    project.embedding_model = "test-model"

    # Create module and record
    module = await create_checklist_module(db_session, project_id=project.id)
    user = await create_user(db_session)
    actor = authenticated(user)

    await create_mock_data_record(
        db_session, module_id=module.id, fields={"name": "TestService"}, created_by=user.id
    )
    await db_session.commit()

    # Create service with max_rows=5
    settings = Settings(mock_data_export_max_rows=5)
    service = MockDataDatasetService(db_session, settings)

    # Should return valid JSON bytes
    result = await service.export_json(module.id, actor=actor)
    payload = json.loads(result)
    assert isinstance(payload, list)
    assert len(payload) == 1
    assert payload[0]["name"] == "TestService"


async def test_export_xlsx_succeeds_under_cap(db_session: AsyncSession) -> None:
    """export_xlsx returns valid workbook when records are under cap."""
    # Create project with embedding_collection
    project = await create_project(db_session)
    project.embedding_collection = "col"
    project.embedding_model = "test-model"

    # Create module and records
    module = await create_checklist_module(db_session, project_id=project.id)
    user = await create_user(db_session)
    actor = authenticated(user)

    # Create records with explicit timestamps to control ordering
    started = datetime.now(UTC)
    acme = MockDataRecord(
        id=uuid.uuid4(),
        checklist_module_id=module.id,
        fields={"name": "Acme"},
        created_by=user.id,
        created_at=started,
    )
    globex = MockDataRecord(
        id=uuid.uuid4(),
        checklist_module_id=module.id,
        fields={"name": "Globex", "start": "2026-01-01"},
        created_by=user.id,
        created_at=started + timedelta(seconds=1),
    )
    db_session.add_all([acme, globex])
    await db_session.commit()

    # Create service with max_rows=10
    settings = Settings(mock_data_export_max_rows=10)
    service = MockDataDatasetService(db_session, settings)

    # Should return valid workbook bytes
    result = await service.export_xlsx(module.id, actor=actor)
    book = load_workbook(BytesIO(result))
    sheet = book.active
    header = [cell.value for cell in sheet[1]]
    assert header == ["name", "start"]
    assert [cell.value for cell in sheet[2]] == ["Acme", None]
    assert [cell.value for cell in sheet[3]] == ["Globex", "2026-01-01"]


async def test_export_refuses_over_the_row_cap(db_session: AsyncSession) -> None:
    """Service rejects export when record count exceeds cap."""
    # Create project with embedding_collection
    project = await create_project(db_session)
    project.embedding_collection = "col"
    project.embedding_model = "test-model"

    # Create module and records with explicit timestamps
    module = await create_checklist_module(db_session, project_id=project.id)
    user = await create_user(db_session)
    actor = authenticated(user)

    started = datetime.now(UTC)
    records = []
    for i in range(2):
        record = MockDataRecord(
            id=uuid.uuid4(),
            checklist_module_id=module.id,
            fields={"id": f"r{i + 1}", "name": "Test"},
            created_by=user.id,
            created_at=started + timedelta(seconds=i),
        )
        records.append(record)
    db_session.add_all(records)
    await db_session.commit()

    # Create service with max_rows=1
    settings = Settings(mock_data_export_max_rows=1)
    service = MockDataDatasetService(db_session, settings)

    # Should raise 409 when export exceeds cap
    with pytest.raises(AppError) as excinfo:
        await service.export_json(module.id, actor=actor)
    assert excinfo.value.status_code == status.HTTP_409_CONFLICT


async def test_export_json_404s_on_nonexistent_module(db_session: AsyncSession) -> None:
    """export_json raises 404 for nonexistent module id."""
    service = MockDataDatasetService(db_session, Settings())
    user = await create_user(db_session)

    with pytest.raises(AppError) as excinfo:
        await service.export_json(uuid.uuid4(), actor=authenticated(user))

    assert excinfo.value.status_code == status.HTTP_404_NOT_FOUND


async def test_export_xlsx_404s_on_nonexistent_module(db_session: AsyncSession) -> None:
    """export_xlsx raises 404 for nonexistent module id."""
    service = MockDataDatasetService(db_session, Settings())
    user = await create_user(db_session)

    with pytest.raises(AppError) as excinfo:
        await service.export_xlsx(uuid.uuid4(), actor=authenticated(user))

    assert excinfo.value.status_code == status.HTTP_404_NOT_FOUND
